#!/usr/bin/env python3
"""按明确区块格式化复杂 Excel，不猜测表格结构，不改动单元格值。"""

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from make_excel import (
    BOTTOM_BORDER,
    FMT,
    HDR_ALIGN,
    LEFT_ALIGN,
    MID_BORDER,
    RIGHT_ALIGN,
    ROW_HEIGHT,
    SUM_KEYWORDS,
    TOP_BORDER,
    _build_theme_styles,
    _backup_for_in_place_edit,
    _cell_display_text,
    _estimate_col_width_from_cells,
    _infer_col_type_from_ws,
    _sample_ws_column,
    _validate_freeze_rows,
)


必需字段 = ("header_row", "data_end_row", "start_col", "end_col")


def _validate_sections(sections):
    """校验区块坐标，并拒绝相互覆盖的矩形。"""
    if not sections:
        raise ValueError("至少提供一个格式化区块")
    normalized = []
    for index, section in enumerate(sections, start=1):
        missing = [key for key in 必需字段 if key not in section]
        if missing:
            raise ValueError(f"第 {index} 个区块缺少字段: {missing}")
        values = tuple(section[key] for key in 必需字段)
        if any(isinstance(value, bool) or not isinstance(value, int) for value in values):
            raise ValueError(f"第 {index} 个区块的坐标必须是整数")
        header_row, data_end_row, start_col, end_col = values
        if header_row < 1 or data_end_row < header_row:
            raise ValueError(f"第 {index} 个区块的行范围无效")
        if start_col < 1 or end_col < start_col:
            raise ValueError(f"第 {index} 个区块的列范围无效")
        normalized.append(dict(zip(必需字段, values)))

    for left_index, left in enumerate(normalized):
        for right in normalized[left_index + 1:]:
            rows_overlap = not (
                left["data_end_row"] < right["header_row"]
                or right["data_end_row"] < left["header_row"]
            )
            columns_overlap = not (
                left["end_col"] < right["start_col"]
                or right["end_col"] < left["start_col"]
            )
            if rows_overlap and columns_overlap:
                raise ValueError("格式化区块不能重叠")
    return normalized


def format_worksheet_sections(
    worksheet,
    sections,
    *,
    theme="default",
    freeze_rows=None,
):
    """格式化单个工作表中明确指定的多个矩形区块。"""
    styles = _build_theme_styles(theme)
    _validate_freeze_rows(freeze_rows)
    sections = _validate_sections(sections)

    for section in sections:
        header_row = section["header_row"]
        data_end_row = section["data_end_row"]
        start_col = section["start_col"]
        end_col = section["end_col"]

        for row in range(header_row, data_end_row + 1):
            worksheet.row_dimensions[row].height = ROW_HEIGHT

        for col in range(start_col, end_col + 1):
            header = worksheet.cell(header_row, col)
            header.font = styles["header_font"]
            header.fill = styles["header_fill"]
            header.alignment = HDR_ALIGN

            samples = _sample_ws_column(
                worksheet,
                col,
                header_row + 1,
                data_end_row,
            )
            column_type = _infer_col_type_from_ws(header.value, samples)
            numeric = column_type in ("money", "number", "pct")
            number_format = FMT.get(column_type)
            alignment = RIGHT_ALIGN if numeric else LEFT_ALIGN
            for row in range(header_row + 1, data_end_row + 1):
                cell = worksheet.cell(row, col)
                cell.alignment = alignment
                if number_format:
                    cell.number_format = number_format
                if numeric:
                    is_formula = (
                        _cell_display_text(cell).startswith("=")
                        or cell.data_type == "f"
                    )
                    cell.font = (
                        styles["data_font_black"]
                        if is_formula
                        else styles["data_font_blue"]
                    )
                else:
                    cell.font = Font(name="Arial", size=11)

            column_letter = get_column_letter(col)
            estimated = _estimate_col_width_from_cells(
                worksheet,
                col,
                header_row,
                data_end_row,
            )
            worksheet.column_dimensions[column_letter].width = max(
                worksheet.column_dimensions[column_letter].width or 0,
                estimated,
            )

        last_label = next(
            (
                worksheet.cell(data_end_row, col).value
                for col in range(start_col, end_col + 1)
                if worksheet.cell(data_end_row, col).value is not None
            ),
            None,
        )
        if last_label is not None and any(
            keyword in str(last_label) for keyword in SUM_KEYWORDS
        ):
            for col in range(start_col, end_col + 1):
                cell = worksheet.cell(data_end_row, col)
                cell.font = styles["sum_font"]
                cell.fill = styles["sum_fill"]

        for row in range(header_row, data_end_row + 1):
            border = (
                TOP_BORDER
                if row == header_row
                else BOTTOM_BORDER
                if row == data_end_row
                else MID_BORDER
            )
            for col in range(start_col, end_col + 1):
                worksheet.cell(row, col).border = border

    worksheet.sheet_view.showGridLines = False
    if freeze_rows is not None:
        worksheet.freeze_panes = None if freeze_rows == 0 else f"A{freeze_rows + 1}"
    return worksheet


def format_excel_sections(
    input_path,
    sections,
    output_path=None,
    *,
    theme="default",
    freeze_rows=None,
    backup=True,
):
    """打开工作簿，格式化活动工作表的明确区块并保存。"""
    source = Path(input_path)
    if not source.is_file():
        raise FileNotFoundError(f"找不到输入文件: {source}")

    destination = Path(output_path) if output_path else source
    if backup and output_path is None:
        _backup_for_in_place_edit(source)

    workbook = load_workbook(source)
    format_worksheet_sections(
        workbook.active,
        sections,
        theme=theme,
        freeze_rows=freeze_rows,
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return str(destination.absolute())


def _parse_section(value):
    try:
        header_row, data_end_row, start_col, end_col = map(int, value.split(":"))
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            "区块格式必须是 表头行:结束行:起始列号:结束列号"
        ) from exc
    return {
        "header_row": header_row,
        "data_end_row": data_end_row,
        "start_col": start_col,
        "end_col": end_col,
    }


def main():
    parser = argparse.ArgumentParser(description="格式化复杂 Excel 的明确区块")
    parser.add_argument("input", help="输入 xlsx")
    parser.add_argument("output", nargs="?", help="输出 xlsx；不传则原地处理")
    parser.add_argument(
        "--section",
        action="append",
        required=True,
        type=_parse_section,
        help="可重复：表头行:结束行:起始列号:结束列号",
    )
    parser.add_argument("--theme", default="default", help="主题名")
    parser.add_argument("--freeze-rows", type=int, default=None, help="冻结行数")
    args = parser.parse_args()
    result = format_excel_sections(
        args.input,
        args.section,
        args.output,
        theme=args.theme,
        freeze_rows=args.freeze_rows,
    )
    print(f"已格式化: {result}")


if __name__ == "__main__":
    main()
