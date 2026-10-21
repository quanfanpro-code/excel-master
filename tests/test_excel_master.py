import sys
import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula


脚本目录 = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(脚本目录))

import make_excel as 被测模块


class ExcelMaster核心测试(unittest.TestCase):
    def test_包含十二套主题并可构建样式(self):
        预期主题 = {
            "default",
            "deep-navy",
            "jade",
            "slate",
            "burgundy",
            "coral",
            "sakura",
            "warm-sun",
            "lavender",
            "matcha",
            "peach",
            "misty",
        }
        self.assertEqual(预期主题, set(被测模块.THEMES))
        for 主题 in 预期主题:
            self.assertIn("header_fill", 被测模块._build_theme_styles(主题))

    def test_小数金额不误判且本地扩展关键词保留(self):
        小数数据 = pd.Series([0.84, 0.71, 0.63])
        self.assertEqual("money", 被测模块._infer_column_type("字段A", 小数数据))
        self.assertEqual("pct", 被测模块._infer_column_type("毛利率", 小数数据))
        self.assertEqual("pct", 被测模块._infer_column_type("gross margin", 小数数据))
        self.assertEqual("money", 被测模块._infer_column_type("返点", 小数数据))
        self.assertEqual("money", 被测模块._infer_column_type("commission", 小数数据))

    def test_安全显示数字日期和公式对象(self):
        工作簿 = Workbook()
        工作表 = 工作簿.active
        工作表["A1"] = 123.45
        工作表["A2"] = date(2026, 7, 26)
        工作表["A3"] = "=SUM(B1:B2)"
        工作表["A4"] = ArrayFormula("A4:A5", "=SUM(B4:B5)")
        工作表["A6"] = DataTableFormula("A6:A7")

        self.assertEqual("123.45", 被测模块._cell_display_text(工作表["A1"]))
        self.assertEqual("2026-07-26", 被测模块._cell_display_text(工作表["A2"]))
        self.assertEqual("=SUM(B1:B2)", 被测模块._cell_display_text(工作表["A3"]))
        self.assertEqual("=SUM(B4:B5)", 被测模块._cell_display_text(工作表["A4"]))
        self.assertEqual("", 被测模块._cell_display_text(工作表["A6"]))

    def test_A列表头为空但数据存在时保持可见且空表头不填色(self):
        工作簿 = Workbook()
        工作表 = 工作簿.active
        工作表["B1"] = "金额"
        工作表["A2"] = "部门甲"
        工作表["B2"] = 100

        被测模块._beautify_worksheet(工作表)

        self.assertGreaterEqual(工作表.column_dimensions["A"].width, 被测模块.MIN_COL_WIDTH)
        self.assertIsNone(工作表["A1"].fill.fill_type)

    def test_冻结零行会真正取消已有冻结(self):
        工作簿 = Workbook()
        工作表 = 工作簿.active
        工作表["B1"] = "金额"
        工作表["B2"] = 100
        工作表.freeze_panes = "A2"

        被测模块._beautify_worksheet(工作表, freeze_rows=0)

        self.assertIsNone(工作表.freeze_panes)

    def test_冻结三行定位到第四行(self):
        工作簿 = Workbook()
        工作表 = 工作簿.active
        工作表["B1"] = "金额"
        工作表["B2"] = 100

        被测模块._beautify_worksheet(工作表, freeze_rows=3)

        self.assertEqual("A4", 工作表.freeze_panes)

    def test_条件格式默认不改且显式启用只调整色阶(self):
        def 创建工作表():
            工作簿 = Workbook()
            工作表 = 工作簿.active
            工作表["B1"] = "金额"
            工作表["B2"] = 1
            工作表["B3"] = 2
            工作表.conditional_formatting.add(
                "B2:B3",
                ColorScaleRule(
                    start_type="min",
                    start_color="FFFFFF",
                    end_type="max",
                    end_color="00B050",
                ),
            )
            工作表.conditional_formatting.add(
                "B2:B3",
                DataBarRule(start_type="min", end_type="max", color="638EC6"),
            )
            return 工作表

        def 读取规则(工作表):
            return [规则 for 条件格式 in 工作表.conditional_formatting for 规则 in 条件格式.rules]

        默认工作表 = 创建工作表()
        默认最大色 = next(
            规则.colorScale.color[-1].rgb
            for 规则 in 读取规则(默认工作表)
            if 规则.type == "colorScale"
        )
        被测模块._beautify_worksheet(默认工作表, theme="coral")
        默认规则 = 读取规则(默认工作表)
        self.assertEqual(
            默认最大色,
            next(
                规则.colorScale.color[-1].rgb
                for 规则 in 默认规则
                if 规则.type == "colorScale"
            ),
        )

        适配工作表 = 创建工作表()
        被测模块._beautify_worksheet(
            适配工作表,
            theme="coral",
            adapt_conditional_formatting=True,
        )
        适配规则 = 读取规则(适配工作表)
        self.assertEqual(2, len(适配规则))
        self.assertTrue(any(规则.type == "dataBar" for 规则 in 适配规则))
        最大色 = next(
            规则.colorScale.color[-1].rgb
            for 规则 in 适配规则
            if 规则.type == "colorScale"
        )
        self.assertTrue(最大色.endswith("D84B4B"))

    def test_空数据在保存前给出明确错误(self):
        with patch.object(被测模块.Workbook, "save") as 保存方法:
            with self.assertRaisesRegex(ValueError, "至少包含一个非空"):
                被测模块.make_excel(pd.DataFrame(), "不会写入.xlsx")
        保存方法.assert_not_called()

    def test_美化工作表不改变公式和值(self):
        工作簿 = Workbook()
        工作表 = 工作簿.active
        工作表["B1"] = "金额"
        工作表["B2"] = 100
        工作表["B3"] = "=B2*2"
        原始值 = [工作表["B2"].value, 工作表["B3"].value]

        被测模块._beautify_worksheet(工作表)

        self.assertEqual(原始值, [工作表["B2"].value, 工作表["B3"].value])

    def test_大表美化不逐行创建行高记录(self):
        工作表 = Workbook().active
        工作表["B1"] = "金额"
        for 行号 in range(2, 2002):
            工作表.cell(行号, 2, 行号)
        被测模块._beautify_worksheet(工作表)
        self.assertEqual(被测模块.ROW_HEIGHT, 工作表.sheet_format.defaultRowHeight)
        self.assertEqual(0, len(工作表.row_dimensions))
        self.assertEqual("medium", 工作表["B2001"].border.bottom.style)

    def test_稀疏表范围检测不填充空单元格(self):
        工作表 = Workbook().active
        工作表["B1"] = "金额"
        工作表["B10000"] = 1
        self.assertEqual(2, len(工作表._cells))
        self.assertEqual(10000, 被测模块._detect_data_range(工作表)[2])
        self.assertEqual(2, len(工作表._cells))

    def test_原地美化前备份到用户统一备份目录(self):
        with TemporaryDirectory() as 临时目录:
            临时根目录 = Path(临时目录)
            源文件 = 临时根目录 / "原始报表.xlsx"
            工作簿 = Workbook()
            工作簿.active["B1"] = "金额"
            工作簿.active["B2"] = 100
            工作簿.save(源文件)
            原始内容 = 源文件.read_bytes()

            with patch.object(被测模块.Path, "home", return_value=临时根目录):
                被测模块.beautify(str(源文件))

            备份文件 = list((临时根目录 / "BackUp").glob("excel-master_*/*"))
            self.assertEqual(1, len(备份文件))
            self.assertEqual(源文件.name, 备份文件[0].name)
            self.assertEqual(原始内容, 备份文件[0].read_bytes())
            self.assertEqual([], list(临时根目录.glob("*_backup_*.xlsx")))


if __name__ == "__main__":
    unittest.main()
