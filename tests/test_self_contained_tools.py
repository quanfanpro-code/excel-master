import sys
import unittest
from pathlib import Path

from openpyxl import Workbook


脚本目录 = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(脚本目录))


class 自包含工具测试(unittest.TestCase):
    def test_显式区块格式化保留值并应用主题(self):
        import format_excel_sections as 区块模块

        工作簿 = Workbook()
        工作表 = 工作簿.active
        工作表["B3"] = "项目"
        工作表["C3"] = "金额"
        工作表["B4"] = "收入"
        工作表["C4"] = 100
        工作表["B5"] = "合计"
        工作表["C5"] = "=SUM(C4:C4)"
        原始值 = [工作表["B4"].value, 工作表["C4"].value, 工作表["C5"].value]

        区块模块.format_worksheet_sections(
            工作表,
            [
                {
                    "header_row": 3,
                    "data_end_row": 5,
                    "start_col": 2,
                    "end_col": 3,
                }
            ],
            theme="coral",
            freeze_rows=3,
        )

        self.assertEqual("solid", 工作表["B3"].fill.fill_type)
        self.assertTrue(工作表["B3"].fill.start_color.rgb.endswith("D84B4B"))
        self.assertEqual("medium", 工作表["B3"].border.top.style)
        self.assertEqual("medium", 工作表["C5"].border.bottom.style)
        self.assertEqual("A4", 工作表.freeze_panes)
        self.assertEqual(原始值, [工作表["B4"].value, 工作表["C4"].value, 工作表["C5"].value])

    def test_区块重叠会被拒绝(self):
        import format_excel_sections as 区块模块

        工作簿 = Workbook()
        工作表 = 工作簿.active
        重叠区块 = [
            {"header_row": 2, "data_end_row": 6, "start_col": 2, "end_col": 5},
            {"header_row": 5, "data_end_row": 8, "start_col": 4, "end_col": 7},
        ]

        with self.assertRaisesRegex(ValueError, "重叠"):
            区块模块.format_worksheet_sections(工作表, 重叠区块)

    def test_整表区块不逐行创建行高记录(self):
        import format_excel_sections as 区块模块

        工作表 = Workbook().active
        工作表["B1"] = "金额"
        for 行号 in range(2, 2002):
            工作表.cell(行号, 2, 行号)
        区块模块.format_worksheet_sections(
            工作表,
            [{"header_row": 1, "data_end_row": 2001, "start_col": 2, "end_col": 2}],
        )
        self.assertEqual(18, 工作表.sheet_format.defaultRowHeight)
        self.assertEqual(0, len(工作表.row_dimensions))
        self.assertEqual("medium", 工作表["B2001"].border.bottom.style)


if __name__ == "__main__":
    unittest.main()
